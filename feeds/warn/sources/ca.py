"""California EDD: the current fiscal year's spreadsheet only.

Older years are PDFs we never need: the feed keeps its own rolling window, so
notices survive the July 1 rollover when EDD resets the spreadsheet.
"""

import io
import re

from openpyxl import load_workbook

from vibefeed import clean, parse_date, parse_int

from .common import city_from_address, kind_of

CODE = "CA"
NAME = "California"
URL = "https://edd.ca.gov/en/jobs_and_training/Layoff_Services_WARN"


def fetch(http, since):
    html = http.get(URL).text
    links = sorted(set(re.findall(r'href="([^"]*warn[^"]*\.xlsx)"', html, re.I)))
    if not links:
        raise RuntimeError("no WARN spreadsheet link on the EDD page")
    rows = []
    for href in links:
        url = href if href.startswith("http") else "https://edd.ca.gov" + href
        wb = load_workbook(io.BytesIO(http.get(url).content), read_only=True, data_only=True)
        name = next((n for n in wb.sheetnames if n.strip().lower() == "detailed warn report"), None)
        ws = wb[name] if name else wb.worksheets[0]
        header = None
        for r in ws.iter_rows(values_only=True):
            first = clean(r[0]).lower() if r and r[0] else ""
            if header is None:
                if first.startswith("county"):
                    header = [clean(c).lower() for c in r]
                    col = lambda *keys: next(i for i, h in enumerate(header) if all(k in h for k in keys))
                    i_notice, i_proc, i_eff = col("notice"), col("processed"), col("effective")
                    i_co, i_type, i_n, i_addr = col("company"), col("layoff"), col("employees"), col("address")
                    i_ind = next((i for i, h in enumerate(header) if "industry" in h), None)
                continue
            if not first or first == "report summary":
                if first == "report summary":
                    break
                continue
            kind, tmp = kind_of(r[i_type])
            addr = clean(r[i_addr])
            rows.append({
                "co": clean(r[i_co]),
                "loc": city_from_address(addr, "CA") or clean(r[0]).removesuffix(" County"),
                "n": parse_int(r[i_n]),
                "kind": kind,
                "tmp": tmp,
                "posted": parse_date(r[i_proc]) or parse_date(r[i_notice]),
                "eff": parse_date(r[i_eff]),
                "ind": re.sub(r"^[\d\s-]+", "", clean(r[i_ind])) if i_ind is not None else "",
                "addr": addr,
            })
        if header is None:
            raise RuntimeError(f"no header row in {url}")
    return rows
