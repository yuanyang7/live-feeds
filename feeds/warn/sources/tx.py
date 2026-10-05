"""Texas Workforce Commission: one spreadsheet per year; we read the latest two."""

import io
import re

from openpyxl import load_workbook

from vibefeed import clean, parse_date, parse_int

CODE = "TX"
NAME = "Texas"
URL = "https://www.twc.texas.gov/data-reports/warn-notice"
# TWC's firewall challenges browser-like UAs (it expects JS) but serves plain
# clients, so identify honestly instead.
HEADERS = {"User-Agent": "live-feeds/1.0 (WARN notice tracker; +https://github.com/yuanyang7/live-feeds)"}


def fetch(http, since):
    html = http.get(URL, headers=HEADERS).text
    links = {}
    for href in re.findall(r'href="(/sites/default/files/[^"]*warn-act-listings-(\d{4})[^"]*\.xlsx)"', html):
        links[int(href[1])] = href[0]
    if not links:
        raise RuntimeError("no WARN spreadsheet links on the TWC page")
    rows = []
    for year in sorted(links)[-2:]:
        wb = load_workbook(io.BytesIO(http.get("https://www.twc.texas.gov" + links[year], headers=HEADERS).content),
                           read_only=True, data_only=True)
        it = wb.worksheets[0].iter_rows(values_only=True)
        header = [clean(h).upper() for h in next(it)]
        col = header.index
        i_notice, i_co, i_county, i_n = col("NOTICE_DATE"), col("JOB_SITE_NAME"), col("COUNTY_NAME"), col("TOTAL_LAYOFF_NUMBER")
        i_eff, i_recv, i_city = col("LAYOFF_DATE"), col("WFDD_RECEIVED_DATE"), col("CITY_NAME")
        for r in it:
            if not r or not r[i_co]:
                continue
            rows.append({
                "co": clean(r[i_co]),
                "loc": clean(r[i_city]) or clean(r[i_county]),
                "n": parse_int(r[i_n]),
                "kind": "",  # TWC doesn't say
                "tmp": False,
                "posted": parse_date(r[i_recv]) or parse_date(r[i_notice]),
                "eff": parse_date(r[i_eff]),
            })
    return rows
