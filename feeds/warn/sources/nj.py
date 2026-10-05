"""New Jersey LWD: one workbook, one sheet per year ("2026 WARN Notices").
NJ only gives the month posted, so posted dates are month-precision ("pm")."""

import io
import re
from datetime import date

from openpyxl import load_workbook

from vibefeed import clean, parse_date, parse_int

CODE = "NJ"
NAME = "New Jersey"
URL = "https://www.nj.gov/labor/employer-services/warn/"
XLSX_URL = "https://www.nj.gov/labor/assets/PDFs/WARN/WARN_Notice_Archive.xlsx"
MONTHS = {m: i for i, m in enumerate(
    "january february march april may june july august september october november december".split(), 1)}


def fetch(http, since):
    wb = load_workbook(io.BytesIO(http.get(XLSX_URL).content), read_only=True, data_only=True)
    sheets = [(int(m.group(1)), ws) for ws in wb.worksheets if (m := re.match(r"(\d{4})", ws.title))]
    if not sheets:
        raise RuntimeError("no year sheets in the NJ workbook")
    rows = []
    for year, ws in sorted(sheets, key=lambda s: s[0])[-2:]:
        it = ws.iter_rows(values_only=True)
        next(it)  # Company, City, Month Posted, Effective Date, Workforce Affected
        for r in it:
            if not r or not clean(r[0]):
                continue
            month = MONTHS.get(clean(r[2]).lower())
            if not month:
                continue
            eff = parse_date(r[3]) or clean(r[3]) or None
            rows.append({
                "co": clean(r[0]),
                "loc": clean(r[1]).title() if clean(r[1]).isupper() else clean(r[1]),
                "n": parse_int(r[4]),
                "kind": "",
                "tmp": False,
                "posted": date(year, month, 1),
                "pm": True,
                "eff": eff,
            })
    return rows
